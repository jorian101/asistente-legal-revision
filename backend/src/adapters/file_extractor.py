"""Adapter de extracción de texto vía PyMuPDF (fitz) y pymupdf4llm.

Dos implementaciones:
- PyMuPdfRawExtractor: fitz.open() + page.get_text("blocks") — preserva
  coordenadas físicas (x0,y0,x1,y1), orden de lectura, y permite distinguir
  headers/footers por posición Y.
- PyMuPdfMarkdownExtractor: pymupdf4llm.to_markdown() — salida markdown
  estructurada (títulos #, listas -, tablas), mejor para segmentación semántica.

Ambos validan magic number PDF (Trail of Bits Regla 3) y límite de tamaño.
"""

from __future__ import annotations

import asyncio
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import fitz  # PyMuPDF

from src.adapters.pdf_worker import SALIDA_VALIDACION
from src.application.ports.text_extractor import (
    ExtractionMode,
    ExtractionResult,
    TextBlock,
    TextExtractor,
)


def _importar_pymupdf4llm():  # type: ignore[no-untyped-def]
    """Import perezoso de pymupdf4llm: arrastra onnxruntime (~1 s y ~100 MB)."""
    try:
        import pymupdf4llm  # type: ignore[import-untyped]
    except Exception as exc:
        raise RuntimeError(
            "pymupdf4llm no está instalado. Añádelo a pyproject.toml o usa modo 'raw'."
        ) from exc
    return pymupdf4llm


# ----- Constantes de validación (Trail of Bits Regla 3) -----
PDF_MAGIC = b"%PDF"
_OCR_TIMEOUT_S = 3600  # tope del subproceso OCR (PDF escaneados grandes)
_PDF_WORKER_TIMEOUT_S = 600  # tope del worker de extraccion nativa (PDF de hasta 500 MB)
_BACKEND_ROOT = Path(__file__).resolve().parents[2]  # cwd del worker (`python -m src...`)


def _comando_worker(file_path: Path, max_size_bytes: int, pages: list[int]) -> list[str]:
    """Comando del subproceso aislado de extraccion nativa (ver pdf_worker.py)."""
    return [
        sys.executable,
        "-m",
        "src.adapters.pdf_worker",
        str(file_path),
        str(max_size_bytes),
        json.dumps(pages),
    ]


MAX_PDF_SIZE_DEFAULT = 50 * 1024 * 1024  # 50 MB (Regla 3 Trail of Bits; configurable via Settings)


class _PyMuPdfBase:
    """Base compartida: validación, apertura, metadatos."""

    def __init__(self, max_size_bytes: int = MAX_PDF_SIZE_DEFAULT):
        self.max_size_bytes = max_size_bytes

    def _validar_archivo(self, file_path: Path) -> None:
        """Valida existencia, tamaño y magic number sin abrir el PDF (no usa fitz)."""
        if not file_path.exists():
            raise FileNotFoundError(f"Archivo no encontrado: {file_path}")

        # Tamaño
        size = file_path.stat().st_size
        if size > self.max_size_bytes:
            raise ValueError(
                f"Archivo excede límite ({size / 1024 / 1024:.1f} MB > "
                f"{self.max_size_bytes / 1024 / 1024:.0f} MB): {file_path}"
            )

        # Magic number (primeros 4 bytes)
        with file_path.open("rb") as f:
            header = f.read(4)
        if header != PDF_MAGIC:
            raise ValueError(
                f"Magic number inválido (esperado %PDF, recibido {header!r}): {file_path}. "
                "Posible archivo corrupto o no-PDF."
            )

    def _validate_and_open(self, file_path: Path) -> fitz.Document:
        """Valida magic number, tamaño y abre documento."""
        self._validar_archivo(file_path)

        # Abrir con PyMuPDF
        doc = fitz.open(file_path)
        if doc.is_encrypted:
            # PDFs encriptados no soportados en esta fase (requiere password)
            doc.close()
            raise ValueError(f"PDF encriptado no soportado: {file_path}")

        return doc

    def _extract_metadata(self, doc: fitz.Document) -> dict:
        """Extrae metadatos básicos del PDF."""
        meta = doc.metadata or {}
        return {
            "title": meta.get("title", ""),
            "author": meta.get("author", ""),
            "subject": meta.get("subject", ""),
            "creator": meta.get("creator", ""),
            "producer": meta.get("producer", ""),
            "creation_date": meta.get("creationDate", ""),
            "modification_date": meta.get("modDate", ""),
            "page_count": doc.page_count,
            "is_encrypted": doc.is_encrypted,
            "is_pdf": True,
        }

    def _pages_to_extract(
        self,
        doc: fitz.Document,
        pages: list[int] | None,
    ) -> list[int]:
        """Normaliza lista de páginas 1-indexed a índices 0-based válidos.

        Si pages es None o lista vacía, se procesan todas las páginas.
        """
        total = doc.page_count
        if not pages:
            return list(range(total))

        result = []
        for p in pages:
            if p < 1 or p > total:
                raise ValueError(f"Página {p} fuera de rango (1-{total})")
            result.append(p - 1)  # a 0-based
        return result


class PyMuPdfRawExtractor(_PyMuPdfBase, TextExtractor):
    """Extractor modo RAW: bloques con coordenadas físicas (fitz.get_text("blocks"))."""

    async def extract(
        self,
        file_path: Path,
        pages: list[int] | None = None,
    ) -> ExtractionResult:
        pages_to_use = pages if pages is not None else []
        return await self.extract_pages(file_path, pages_to_use)

    async def extract_pages(
        self,
        file_path: Path,
        pages: list[int],
    ) -> ExtractionResult:
        # Ejecutar en thread pool para no bloquear event loop (fitz es sync CPU-bound)
        loop = asyncio.get_running_loop()
        return await loop.run_in_executor(None, self._extract_sync, file_path, pages)

    def _extract_sync(
        self,
        file_path: Path,
        pages: list[int],
    ) -> ExtractionResult:
        doc = self._validate_and_open(file_path)
        try:
            page_indices = self._pages_to_extract(doc, pages)
            metadata = self._extract_metadata(doc)

            all_blocks: list[TextBlock] = []
            text_parts: list[str] = []

            for idx in page_indices:
                page = doc[idx]
                # get_text("blocks") devuelve tuplas: (x0, y0, x1, y1, text, block_no, block_type)
                raw_blocks = page.get_text("blocks", sort=True)

                for b in raw_blocks:
                    x0, y0, x1, y1, text, block_no, block_type = b
                    # Filtrar bloques vacíos o solo whitespace
                    if text and text.strip():
                        tb = TextBlock(
                            x0=x0,
                            y0=y0,
                            x1=x1,
                            y1=y1,
                            text=text.strip(),
                            block_no=block_no,
                            block_type=block_type,
                        )
                        all_blocks.append(tb)
                        text_parts.append(text.strip())

            full_text = "\n\n".join(text_parts)
            extracted_pages_1idx = [i + 1 for i in page_indices]

            return ExtractionResult(
                full_text=full_text,
                blocks=all_blocks,
                pages_count=doc.page_count,
                extracted_pages=extracted_pages_1idx,
                metadata=metadata,
            )
        finally:
            doc.close()


class PyMuPdfMarkdownExtractor(_PyMuPdfBase, TextExtractor):
    """Extractor modo MARKDOWN: pymupdf4llm.to_markdown() → markdown estructurado."""

    async def extract(
        self,
        file_path: Path,
        pages: list[int] | None = None,
    ) -> ExtractionResult:
        pages_to_use = pages if pages is not None else []
        return await self.extract_pages(file_path, pages_to_use)

    async def extract_pages(
        self,
        file_path: Path,
        pages: list[int],
    ) -> ExtractionResult:
        _importar_pymupdf4llm()  # falla rapido y claro si falta

        loop = asyncio.get_running_loop()
        return await loop.run_in_executor(None, self._extract_sync, file_path, pages)

    def _extract_sync(
        self,
        file_path: Path,
        pages: list[int],
    ) -> ExtractionResult:
        doc = self._validate_and_open(file_path)
        try:
            page_indices = self._pages_to_extract(doc, pages)
            metadata = self._extract_metadata(doc)

            # pymupdf4llm.to_markdown acepta page_numbers (0-based)
            md_text = _importar_pymupdf4llm().to_markdown(
                doc=doc,
                pages=page_indices,
                page_chunks=False,  # texto único concatenado
                write_images=False,
                embed_images=False,
            )

            # Para compatibilidad con ExtractionResult, no hay bloques con coords
            # pero sí podemos devolver lista vacía
            extracted_pages_1idx = [i + 1 for i in page_indices]

            return ExtractionResult(
                full_text=md_text,
                blocks=[],  # no coords en modo markdown
                pages_count=doc.page_count,
                extracted_pages=extracted_pages_1idx,
                metadata=metadata,
            )
        finally:
            doc.close()


class HybridFileExtractor(PyMuPdfRawExtractor):
    """Extractor híbrido: docx (python-docx) + pdf nativo (fitz) + scan (OCR pipeline).

    Router:
    - .docx → python-docx (preserva negrita/alineación, rápido)
    - .pdf con texto → fitz blocks (rápido, coords)
    - .pdf escaneado (sin texto) → delega a tools/formatos pipeline vía subprocess:
      Docling con OCR Tesseract 5.3.4 spa (`TesseractCliOcrOptions` en
      tools/formatos/formatos/extract_scan.py). Se invoca con `--sin-consenso`, así que la
      pasada de PaddleOCR no corre. Si el pipeline no está disponible (sin venv), falla
      con un error claro.

    Pensado para CargarObra: un solo TextExtractor que sube bien tanto
    nativos como escaneados CamScanner.
    """

    def _is_docx(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".docx"

    def _extract_docx_sync(self, file_path: Path) -> ExtractionResult:
        try:
            import docx  # type: ignore[import-untyped]
        except ImportError as exc:
            raise RuntimeError("python-docx no instalado para .docx") from exc

        doc = docx.Document(str(file_path))
        blocks: list[TextBlock] = []
        text_parts: list[str] = []
        for idx, para in enumerate(doc.paragraphs):
            txt = para.text.strip()
            if not txt:
                continue
            # coordenadas ficticias para docx (no hay bbox real)
            tb = TextBlock(
                x0=0,
                y0=idx * 10,
                x1=500,
                y1=idx * 10 + 10,
                text=txt,
                block_no=idx,
                block_type=0,
            )
            blocks.append(tb)
            text_parts.append(txt)
        # tablas
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    txt = cell.text.strip()
                    if txt:
                        text_parts.append(txt)
        full_text = "\n\n".join(text_parts)
        return ExtractionResult(
            full_text=full_text,
            blocks=blocks,
            pages_count=1,
            extracted_pages=[1],
            metadata={"is_pdf": False, "is_docx": True, "page_count": 1},
        )

    def _is_scan_pdf(self, file_path: Path) -> bool:
        """Heurística: pdf sin texto → scan. Usa fitz rápido sin OCR."""
        try:
            doc = fitz.open(file_path)
            sample = min(3, doc.page_count)
            chars = sum(len(doc[i].get_text("text").strip()) for i in range(sample))
            doc.close()
            avg = chars / max(sample, 1)
            return avg < 40
        except Exception:
            return True

    def _extract_via_ocr_pipeline(self, file_path: Path) -> ExtractionResult:
        """Delega a tools/formatos pipeline (Docling + OCR Tesseract 5.3.4 spa) vía subprocess.

        Motor configurado en extract_scan.py (`TesseractCliOcrOptions`); `--sin-consenso`
        omite la pasada de PaddleOCR. El layout trae la página de cada bloque, pero acá se
        descarta (scripts/carga/leer_ocr.py la conserva). Para archivos grandes, el pipeline
        externo ya maneja paginación interna; acá un solo subprocess
        con timeout largo es suficiente y evita el bug de chunk-PDF
        que colgaba en fitz.insert_pdf de scans CamScanner.
        """
        import os

        tools_root = Path(__file__).resolve().parents[3] / "tools" / "formatos"
        if not (tools_root / "formatos" / "cli.py").exists():
            raise RuntimeError(f"Pipeline formatos no encontrado en {tools_root}")

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out"
            venv_python = tools_root / ".venv" / "bin" / "python"
            python_exe = str(venv_python) if venv_python.exists() else sys.executable
            cmd = [
                python_exe,
                "-m",
                "formatos.cli",
                str(file_path),
                "--out",
                str(out),
                "--tipo",
                "otro",
                "--autor",
                "desconocido",
                "--sin-consenso",
            ]
            env = os.environ.copy()
            env["PYTHONPATH"] = str(tools_root) + ":" + env.get("PYTHONPATH", "")
            try:
                subprocess.run(
                    cmd,
                    cwd=str(tools_root),
                    env=env,
                    check=True,
                    capture_output=True,
                    timeout=_OCR_TIMEOUT_S,
                )
            except subprocess.TimeoutExpired as exc:
                raise RuntimeError(
                    f"OCR pipeline excedió el tiempo máximo ({_OCR_TIMEOUT_S} s)"
                ) from exc
            except subprocess.CalledProcessError as exc:
                detalle = (exc.stderr or b"").decode("utf-8", errors="replace")[:500]
                raise RuntimeError(f"OCR pipeline falló: {detalle}") from exc
            except FileNotFoundError as exc:
                raise RuntimeError(f"OCR pipeline no disponible: {exc}") from exc

            layouts = list(out.rglob("layout.json"))
            if not layouts:
                raise RuntimeError("OCR pipeline no generó layout.json")
            data = json.loads(layouts[0].read_text(encoding="utf-8"))
            blocks: list[TextBlock] = []
            text_parts: list[str] = []
            for b in data.get("blocks", []):
                txt = " ".join(r.get("text", "") for r in b.get("runs", []))
                if txt.strip():
                    text_parts.append(txt.strip())
                    bbox = b.get("bbox") or [0, 0, 500, 10]
                    tb = TextBlock(
                        x0=bbox[0],
                        y0=bbox[1],
                        x1=bbox[2],
                        y1=bbox[3],
                        text=txt.strip(),
                        block_no=b.get("index", 0),
                        block_type=0,
                    )
                    blocks.append(tb)
            full_text = "\n\n".join(text_parts)
            # Control de calidad: si hay demasiados caracteres malos, marcar para revisión manual
            # (sin reintento con otro motor, solo flag)
            bad_ratio = 0
            if full_text:
                # Ratio de caracteres no alfanuméricos no comunes (excluye puntuación legal común)
                allowed_punct = set(" .,;:()\"-'°/º\n")
                total = len(full_text)
                bad = sum(
                    1
                    for c in full_text
                    if c not in allowed_punct and not c.isalnum() and not c.isspace()
                )
                bad_ratio = bad / max(total, 1)
            needs_review = bad_ratio > 0.05 or "�" in full_text or len(full_text.strip()) < 100
            return ExtractionResult(
                full_text=full_text,
                blocks=blocks,
                pages_count=data.get("meta", {}).get("pages") or 1,
                extracted_pages=list(range(1, (data.get("meta", {}).get("pages") or 1) + 1)),
                metadata={
                    "is_pdf": True,
                    "is_docx": False,
                    "ocr": True,
                    "engine": data.get("meta", {}).get("engine", "docling"),
                    "bad_ratio": round(bad_ratio, 3),
                    "needs_review": needs_review,
                },
            )

    async def extract(
        self,
        file_path: Path,
        pages: list[int] | None = None,
    ) -> ExtractionResult:
        # docx directo
        if self._is_docx(file_path):
            return await asyncio.to_thread(self._extract_docx_sync, file_path)
        # pdf: chequeos baratos aqui (sin fitz) y todo lo que toca PyMuPDF en un subproceso
        # aislado: un PDF hostil que tumbe la libreria nativa solo mata al worker (R5).
        self._validar_archivo(file_path)
        resultado = await self._extraer_pdf_aislado(file_path, pages or [])
        if resultado is None:  # escaneo (sin texto): pipeline OCR, tambien en subproceso
            return await asyncio.to_thread(self._extract_via_ocr_pipeline, file_path)
        return resultado

    def _extraer_nativo_sync(self, file_path: Path, pages: list[int]) -> ExtractionResult | None:
        """Extraccion nativa con fitz; None si el PDF es un escaneo. Corre en el worker."""
        doc = self._validate_and_open(file_path)
        doc.close()
        if self._is_scan_pdf(file_path):
            return None
        return PyMuPdfRawExtractor._extract_sync(self, file_path, pages)

    async def _extraer_pdf_aislado(
        self, file_path: Path, pages: list[int]
    ) -> ExtractionResult | None:
        """Ejecuta pdf_worker en un subproceso con timeout; None si el PDF es un escaneo."""
        proc = await asyncio.create_subprocess_exec(
            *_comando_worker(file_path, self.max_size_bytes, pages),
            cwd=_BACKEND_ROOT,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        try:
            stdout, stderr = await asyncio.wait_for(
                proc.communicate(), timeout=_PDF_WORKER_TIMEOUT_S
            )
        except TimeoutError:
            raise RuntimeError(
                f"El worker de extracción excedió el tiempo máximo ({_PDF_WORKER_TIMEOUT_S} s)"
            ) from None
        finally:
            if proc.returncode is None:  # timeout o cancelacion (cliente desconectado)
                proc.kill()
                await proc.wait()

        detalle = stderr.decode("utf-8", errors="replace").strip()[:500]
        if proc.returncode == SALIDA_VALIDACION:
            raise ValueError(detalle)
        if proc.returncode != 0:
            raise RuntimeError(
                f"El worker de extracción falló (código {proc.returncode}): {detalle}"
            )
        try:
            datos = json.loads(stdout)
        except ValueError as exc:
            raise RuntimeError("El worker de extracción devolvió una salida inválida") from exc
        if datos["escaneo"]:
            return None
        r = datos["resultado"]
        return ExtractionResult(
            full_text=r["full_text"],
            blocks=[TextBlock(**b) for b in r["blocks"]],
            pages_count=r["pages_count"],
            extracted_pages=r["extracted_pages"],
            metadata=r["metadata"],
        )

    async def extract_pages(
        self,
        file_path: Path,
        pages: list[int],
    ) -> ExtractionResult:
        # para compatibilidad, delega a extract (ignora pages para docx/scan)
        return await self.extract(file_path, pages)

    def _extract_sync(
        self,
        file_path: Path,
        pages: list[int],
    ) -> ExtractionResult:
        # sync helper para compatibilidad con base
        if self._is_docx(file_path):
            return self._extract_docx_sync(file_path)
        if self._is_scan_pdf(file_path):
            return self._extract_via_ocr_pipeline(file_path)
        # fallback a base fitz
        return super()._extract_sync(file_path, pages)  # type: ignore[attr-defined]


# ----- Factory pública -----


def create_text_extractor(mode: str = ExtractionMode.RAW, **kwargs) -> TextExtractor:
    """Factory del extractor según modo.

    Args:
        mode: "raw" (fitz blocks), "markdown" (pymupdf4llm) o "auto" (híbrido).
        **kwargs: pasados al constructor (ej. max_size_bytes).

    Returns:
        Instancia de TextExtractor.

    Raises:
        ValueError: modo desconocido.
    """
    if mode == ExtractionMode.RAW:
        return PyMuPdfRawExtractor(**kwargs)
    elif mode == ExtractionMode.MARKDOWN:
        return PyMuPdfMarkdownExtractor(**kwargs)
    elif mode == ExtractionMode.AUTO:
        return HybridFileExtractor(**kwargs)
    else:
        raise ValueError(
            f"Modo de extracción desconocido: {mode}. "
            f"Valores válidos: {ExtractionMode.RAW}, "
            f"{ExtractionMode.MARKDOWN}, {ExtractionMode.AUTO}"
        )


# ----- Helpers para scripts de inspección -----


async def inspect_file(
    file_path: Path,
    pages: list[int],
    mode: str = ExtractionMode.RAW,
    max_size_bytes: int = MAX_PDF_SIZE_DEFAULT,
) -> ExtractionResult:
    """Helper simple para scripts: extrae y devuelve resultado."""
    extractor = create_text_extractor(mode, max_size_bytes=max_size_bytes)
    return await extractor.extract_pages(file_path, pages)
